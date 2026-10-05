import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

TEMPLATES_DIR = Path(r"c:\Daniel Olate\Soft Berardi\SoftBerardi\app\static\templates")

HEADER_FILL = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
NOTE_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
NOTE_FONT = Font(color="92400E", size=10, italic=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style_header(ws, ncols, nrows=1, start_row=1):
    for col in range(1, ncols + 1):
        for row in range(start_row, start_row + nrows):
            cell = ws.cell(row=row, column=col)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = CENTER
            cell.border = THIN_BORDER


def style_rows(ws, start_row, end_row, ncols):
    for row in range(start_row, end_row + 1):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = THIN_BORDER
            cell.alignment = LEFT_WRAP


def autosize(ws, columns_widths=None, min_width=12, max_width=42):
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        col_letter = get_column_letter(col_idx)
        if columns_widths and col_letter in columns_widths:
            ws.column_dimensions[col_letter].width = columns_widths[col_letter]
            continue
        max_length = min_width
        for cell in column_cells:
            val = "" if cell.value is None else str(cell.value)
            for line in val.split("\n"):
                if len(line) > max_length:
                    max_length = len(line)
        ws.column_dimensions[col_letter].width = min(max_length + 2, max_width)


def build_technicians_xlsx(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tecnicos"

    headers = [
        "employee_code",
        "name",
        "region",
        "phone",
        "commune",
        "team",
        "centro",
        "empresa",
        "sindicato",
        "supervisor",
        "patente",
        "grupo_sanguineo",
        "art",
        "numero_emergencia",
        "alergias",
        "activo",
        "movil",
    ]

    header_comments = {
        "employee_code": "Legajo UNICO por tecnico. Si ya existe, actualiza los datos.\nSi se deja vacio, se intentara usar 'movil' como clave alternativa.",
        "name": "Nombre completo del tecnico. Obligatorio si es un alta nueva.",
        "region": "Region geografica general (ej: Mendoza, Valparaiso, Cuyo, etc.).",
        "phone": "Telefono personal / celular del tecnico.",
        "commune": "Comuna o barrio donde vive/trabaja (ej: Maipu, Godoy Cruz).",
        "team": "Equipo, zona, turno o nombre de cuadrilla.",
        "centro": "Centro / sucursal / codigo de centro (ej: MA01, VP02).",
        "empresa": "Nombre de la empresa contratista.",
        "sindicato": "Sindicato o gremio (ej: UOM, SOEME, UPC).",
        "supervisor": "Nombre completo del supervisor. Si no existe, se crea automaticamente.",
        "patente": "Patente del vehiculo ASIGNADO. Debe existir en la tabla vehiculos.\nDejar VACIO si no tiene vehiculo asignado.",
        "grupo_sanguineo": "Valores admitidos: A+, A-, B+, B-, AB+, AB-, O+, O-\nTambien acepta variantes: A positivo, a-, o NEGATIVO, 0+, etc.",
        "art": "Aseguradora de ART / Aseguradora de riesgos del trabajo (ej: SANCOR SEGUROS, SWISS MEDICAL, OSDE).",
        "numero_emergencia": "Telefono de contacto en caso de EMERGENCIA.",
        "alergias": "Alergias, enfermedades preexistentes o datos medicos relevantes.\nDejar vacio si no tiene nada para informar.",
        "activo": "1 = Activo (asignable a tareas) | 0 = Inactivo (solo historico).\nSi se deja vacio se asume 1.",
        "movil": "Codigo del movil tecnico / numero de equipo (ej: 606). Si no existe se crea automaticamente.",
    }

    examples = [
        ["79", "Juan Perez Mendoza", "Mendoza", "+54 9 261 555-0101", "Maipu", "Zona Norte - Turno Manana", "MA01", "TechConnect S.A.", "UOM", "Alfredo Leiva", "", "O+", "SANCOR SEGUROS", "+54 9 261 555-9999", "Penicilina, mariscos, hipertension", "1", "606"],
        ["606", "Carlos Alberto Mena", "Valparaiso", "+56 9 1111-2222", "Valparaiso", "Cuadrilla Centro", "VP02", "Servicios Integrales Ltda.", "SOEME", "Marta Fuentes", "", "A+", "SWISS MEDICAL", "+56 9 2222-8888", "Asma, alergia a abejas", "1", "606"],
        ["600", "Pedro Soto", "Quilpue", "+56 9 2222-3333", "Quilpue", "Cuadrilla Sur - Turno Tarde", "VP03", "TecnoRed S.A.", "UPC", "Felipe Rojas", "", "B-", "", "+56 9 3333-7777", "", "1", "600"],
    ]

    ws.append(headers)
    style_header(ws, len(headers))

    for col_idx, header_name in enumerate(headers, start=1):
        if header_name in header_comments:
            cell = ws.cell(row=1, column=col_idx)
            comment = Comment(header_comments[header_name], "SoftBerardi")
            comment.width = 380
            comment.height = max(80, 25 * len(header_comments[header_name].split("\n")))
            cell.comment = comment

    example_start_row = 2
    for ex in examples:
        ws.append(ex)
    style_rows(ws, example_start_row, example_start_row + len(examples) - 1, len(headers))
    ws.freeze_panes = "A2"

    col_widths = {
        "A": 16, "B": 28, "C": 16, "D": 22, "E": 16, "F": 26, "G": 10,
        "H": 24, "I": 12, "J": 22, "K": 14, "L": 14, "M": 20, "N": 22,
        "O": 34, "P": 10, "Q": 12,
    }
    autosize(ws, col_widths)

    instructions = wb.create_sheet("Instrucciones")
    ins_rows = [
        ["INSTRUCCIONES PARA COMPLETAR ESTA PLANTILLA"],
        [""],
        ["1) ENCABEZADOS (fila 1 AZUL): NO MODIFIQUES los nombres de las columnas. Si los cambias, el importador no reconocera los datos."],
        ["2) AYUDA en encabezados: pasá el mouse por encima de cada celda azul para ver un comentario con la explicacion y formato esperado."],
        ["3) FILAS 2 A 4: Son EJEMPLOS de tecnicos. Borralas ANTES de enviar la plantilla al supervisor para que complete los datos reales."],
        [""],
        ["COMO ACTUALIZAR DATOS DE TECNICOS EXISTENTES:"],
        ["  a. Completa la columna A 'employee_code' con el LEGAJO correcto (es la clave principal)."],
        ["  b. Completa SOLAMENTE las columnas que quieras modificar. LAS COLUMNAS VACIAS SE RESPETAN (no se pisan con datos vacios)."],
        ["  c. Ejemplo: si solo queres cambiar el ART y el numero de emergencia, completa legajo + art + numero_emergencia y dejando el resto vacio."],
        [""],
        ["COMO DAR DE ALTA TECNICOS NUEVOS:"],
        ["  a. Completa 'employee_code' con el legajo nuevo (o al menos el nombre completo + movil)."],
        ["  b. Datos minimos recomendados: legajo, nombre, region, telefono."],
        [""],
        ["VALIDACIONES IMPORTANTES (la app rechaza filas que no cumplan):"],
        ["  - Grupo sanguineo invalido: se rechaza la fila. Valores admitidos: A+, A-, B+, B-, AB+, AB-, O+, O-."],
        ["  - Patente de vehiculo INEXISTENTE en la tabla vehiculos: se rechaza la fila. Primero crea el vehiculo o dejala vacia."],
        [""],
        ["FOTO DE PERFIL:"],
        ["  - Las fotos NO se importan por Excel. Se deben cargar individualmente en el menu:"],
        ["    Administrar Datos → Tecnicos → (buscar el tecnico) → Editar → Foto de perfil."],
    ]
    for row in ins_rows:
        instructions.append(row)
    instructions["A1"].font = Font(bold=True, size=14, color="1E3A5F")
    for r in range(1, len(ins_rows) + 1):
        c = instructions.cell(row=r, column=1)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if r in [1, 6, 10, 14, 17]:
            c.font = Font(bold=True, size=11 if r > 1 else 14, color="1E3A5F")
    instructions.column_dimensions["A"].width = 140

    wb.save(str(path))
    print(f"✅ {path.name} creado OK")


def build_vehicles_xlsx(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Vehiculos"
    headers = [
        "plate",
        "unit_number",
        "brand",
        "model",
        "year",
        "assigned_employee_code",
        "odometer_km",
        "status",
    ]
    header_comments = {
        "plate": "Patente / Dominio / Matricula. UNICO por vehiculo. Si ya existe, actualiza datos.",
        "unit_number": "Numero interno / Unidad (ej: 501, 502, 1, 15).",
        "brand": "Marca (ej: Toyota, Volkswagen, Ford, Renault).",
        "model": "Modelo (ej: Hilux, Amarok, Ranger, Kangoo).",
        "year": "Anio de fabricacion (ej: 2022, 2019).",
        "assigned_employee_code": "Legajo del tecnico ASIGNADO. Debe existir en tabla tecnicos o se puede dejar vacio.",
        "odometer_km": "Kilometraje actual SOLO NUMEROS (ej: 58300).",
        "status": "Estado: activo, en mantenimiento, fuera de servicio, etc.",
    }
    examples = [
        ["AB123CD", "501", "Toyota", "Hilux", "2022", "", "58300", "activo"],
        ["CD456EF", "502", "Volkswagen", "Amarok", "2021", "", "72000", "activo"],
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    for col_idx, header_name in enumerate(headers, start=1):
        if header_name in header_comments:
            cell = ws.cell(row=1, column=col_idx)
            comment = Comment(header_comments[header_name], "SoftBerardi")
            comment.width = 360
            comment.height = 80
            cell.comment = comment
    for ex in examples:
        ws.append(ex)
    style_rows(ws, 2, 1 + len(examples), len(headers))
    ws.freeze_panes = "A2"
    col_widths = {"A": 18, "B": 14, "C": 16, "D": 20, "E": 10, "F": 24, "G": 16, "H": 22}
    autosize(ws, col_widths)
    wb.save(str(path))
    print(f"✅ {path.name} creado OK")


def build_material_stock_xlsx(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock Materiales"
    headers = ["Material", "Codigo", "Descripcion"] + [str(n) for n in range(600, 611)] + ["Total"]
    ws.append(headers)
    style_header(ws, len(headers))
    header_comments_first = {
        "Material": "Nombre o descripcion corta del material.",
        "Codigo": "Codigo SKU interno (opcional).",
        "Descripcion": "Descripcion larga / detalle del material.",
        "Total": "Suma total de unidades (OPCIONAL). Si la dejas vacia no pasa nada.",
    }
    for col_idx, header_name in enumerate(headers, start=1):
        if header_name in header_comments_first:
            cell = ws.cell(row=1, column=col_idx)
            comment = Comment(header_comments_first[header_name], "SoftBerardi")
            comment.width = 320
            comment.height = 70
            cell.comment = comment
        elif header_name.isdigit():
            cell = ws.cell(row=1, column=col_idx)
            comment = Comment(
                f"Cantidad en stock del MOVIL tecnico Nro {header_name}.\n"
                "Si necesitas mas columnas de moviles, agrega una nueva columna\n"
                "con el NUMERO de movil como titulo (ej: 500, 599, 700).",
                "SoftBerardi"
            )
            comment.width = 320
            comment.height = 100
            cell.comment = comment
    examples = [
        ["Cable FTP Cat 6", "MAT-001", "Cable de red exterior 305m", 10, 12, 8, 15, 6, 9, 11, 7, 13, 5, 4, 100],
        ["Conector SC/APC", "MAT-002", "Conector de fibra optica monomodo", 50, 45, 60, 30, 80, 25, 40, 55, 35, 65, 70, 555],
    ]
    for ex in examples:
        ws.append(ex)
    style_rows(ws, 2, 1 + len(examples), len(headers))
    ws.freeze_panes = "D2"
    col_widths = {"A": 28, "B": 16, "C": 40}
    autosize(ws, col_widths, min_width=10, max_width=20)
    wb.save(str(path))
    print(f"✅ {path.name} creado OK")


if __name__ == "__main__":
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    build_technicians_xlsx(TEMPLATES_DIR / "tecnicos_template.xlsx")
    build_vehicles_xlsx(TEMPLATES_DIR / "vehiculos_template.xlsx")
    build_material_stock_xlsx(TEMPLATES_DIR / "material_stock_template.xlsx")
    print("\n🎉 Todas las plantillas XLSX generadas en:", str(TEMPLATES_DIR))
